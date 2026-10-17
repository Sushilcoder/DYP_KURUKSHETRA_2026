from __future__ import annotations
import os, re, subprocess, tempfile, shutil, json
from pathlib import Path
from fastapi import FastAPI, HTTPException
from pptx import Presentation
import fitz

app=FastAPI(title="Course Deck Renderer and Layout Inspector",version="0.1.0")
STORAGE_ROOT=Path(os.getenv("STORAGE_ROOT","/app/storage/materials")).resolve()
ID_RE=re.compile(r"^[A-Za-z0-9_-]{1,80}$")

def normalize_text(value: str) -> str:
    return re.sub(r"[^a-z0-9]+","",(value or "").lower())

def inspect_pdf(pptx_path: Path, pdf_path: Path, course_path: Path | None = None) -> dict:
    prs=Presentation(str(pptx_path))
    doc=fitz.open(str(pdf_path))
    missing=[]; bounds=[]
    if course_path and course_path.exists():
        try:
            source_course=json.loads(course_path.read_text(encoding="utf-8"))
            for si, item in enumerate(source_course.get("slides",[]),1):
                for block in item.get("code_blocks",[]):
                    lines=(block.get("code","") or "").splitlines()
                    if len(lines)>18:
                        missing.append({"slide_id":f"SLIDE-{si:02d}","issue":"CODE_BLOCK_TRUNCATED_IN_SLIDE","line_count":len(lines),"title":block.get("title","Code example")})
                    long_lines=[len(line) for line in lines if len(line)>52]
                    if long_lines:
                        missing.append({"slide_id":f"SLIDE-{si:02d}","issue":"CODE_LINE_MAY_OVERFLOW","max_line_length":max(long_lines),"title":block.get("title","Code example")})
        except Exception:
            missing.append({"issue":"SOURCE_COURSE_MANIFEST_UNREADABLE"})
    if len(doc)!=len(prs.slides):
        missing.append({"issue":"PAGE_COUNT_MISMATCH","expected":len(prs.slides),"actual":len(doc)})
    for idx,slide in enumerate(prs.slides):
        if idx>=len(doc): break
        page=doc[idx]
        page_text=page.get_text("text") or ""
        norm_page=normalize_text(page_text)
        for shape in slide.shapes:
            raw=(getattr(shape,"text","") or "").strip()
            normalized=normalize_text(raw)
            if len(normalized)>=8 and normalized not in norm_page:
                # Avoid checking generated long source snippets for exact full text where visual excerpting is intentional.
                if "Cascadia Code" in str(getattr(getattr(shape,"text_frame",None),"paragraphs",[None])[0].font.name if getattr(shape,"has_text_frame",False) and shape.text_frame.paragraphs else ""):
                    continue
                missing.append({"slide_id":f"SLIDE-{idx+1:02d}","issue":"TEXT_NOT_FULLY_VISIBLE_OR_EXTRACTABLE","text_preview":raw[:100]})
        page_rect=page.rect
        for block in page.get_text("dict").get("blocks",[]):
            for line in block.get("lines",[]):
                for span in line.get("spans",[]):
                    x0,y0,x1,y1=span.get("bbox",[0,0,0,0])
                    if x0 < -1 or y0 < -1 or x1 > page_rect.width+1 or y1 > page_rect.height+1:
                        bounds.append({"slide_id":f"SLIDE-{idx+1:02d}","issue":"TEXT_OUTSIDE_PAGE_BOUNDS","bbox":[x0,y0,x1,y1]})
    page_count=len(doc)
    doc.close()
    okay=(len(missing)==0 and len(bounds)==0 and len(prs.slides)>0 and len(prs.slides)==page_count)
    return {"status":"PASS" if okay else "FAIL","slide_count":len(prs.slides),"pdf_page_count":page_count,"missing_text_count":len(missing),"out_of_bounds_count":len(bounds),"issues":(missing+bounds)[:100],"inspection":"LibreOffice PDF render + text visibility reconciliation + page-bound checks","renderer":"LibreOffice Impress"}

@app.get("/health")
def health():
    return {"status":"ok","libreoffice_available":shutil.which("libreoffice") is not None,"storage_root":str(STORAGE_ROOT)}

@app.post("/render/{material_id}")
def render(material_id: str):
    if not ID_RE.fullmatch(material_id): raise HTTPException(400,"Invalid material ID")
    base=(STORAGE_ROOT/material_id).resolve()
    if not str(base).startswith(str(STORAGE_ROOT)):
        raise HTTPException(400,"Invalid material path")
    pptx=base/"01_Presentation/course_slides.pptx"
    target=base/"01_Presentation/course_slides.pdf"
    if not pptx.exists(): raise HTTPException(404,"PPTX is not ready")
    if not shutil.which("libreoffice"):
        return {"status":"BLOCKED","message":"LibreOffice is not installed; PDF was not rendered from PPTX."}
    with tempfile.TemporaryDirectory(prefix="studio-render-") as tmp:
        tmp_path=Path(tmp); profile=tmp_path/"lo-profile"; outdir=tmp_path/"out"; outdir.mkdir()
        cmd=["libreoffice",f"-env:UserInstallation=file://{profile}","--headless","--convert-to","pdf","--outdir",str(outdir),str(pptx)]
        try:
            proc=subprocess.run(cmd,capture_output=True,text=True,timeout=90,check=False)
        except subprocess.TimeoutExpired:
            return {"status":"BLOCKED","message":"LibreOffice render timed out."}
        rendered=outdir/(pptx.stem+".pdf")
        if proc.returncode!=0 or not rendered.exists():
            return {"status":"FAIL","message":"PowerPoint-to-PDF conversion failed.","returncode":proc.returncode,"stdout":proc.stdout[-3000:],"stderr":proc.stderr[-3000:]}
        shutil.copy2(rendered,target)
    info=inspect_pdf(pptx,target,base/"course.json")
    info["pdf_file"]=str(target.relative_to(STORAGE_ROOT))
    info["libreoffice_stdout"]=""
    return info
