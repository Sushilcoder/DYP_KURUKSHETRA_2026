from __future__ import annotations
import json, os, re, shutil, subprocess, zipfile
from pathlib import Path
from xml.sax.saxutils import escape
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Preformatted, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from .models import CourseContent


def _hex_rgb(value: str) -> RGBColor:
    v=(value or "#1D4ED8").lstrip("#")
    return RGBColor(int(v[0:2],16),int(v[2:4],16),int(v[4:6],16))


def _write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def _make_pptx(course: CourseContent, path: Path, brand: str, primary: str) -> dict:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    accent = _hex_rgb(primary)
    dark = RGBColor(24, 35, 56)
    white = RGBColor(255, 255, 255)
    overflow = []
    for idx, item in enumerate(course.slides, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        # Embed speaker notes in the actual PPTX as well as exporting one notes file per slide.
        # This makes the deck usable in Presenter View and keeps notes part of the deliverable.
        if item.notes.strip():
            slide.notes_slide.notes_text_frame.text = item.notes.strip()
        bg = slide.background.fill
        bg.solid(); bg.fore_color.rgb = white
        band = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.18))
        band.fill.solid(); band.fill.fore_color.rgb=accent; band.line.fill.background()
        # Header
        title_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.45), Inches(12.0), Inches(0.75))
        tf=title_box.text_frame; tf.clear(); tf.word_wrap=True
        p=tf.paragraphs[0]; p.text=item.title; p.font.name="Aptos Display"; p.font.size=Pt(29 if len(item.title)<50 else 24); p.font.bold=True; p.font.color.rgb=dark
        y=1.42
        left=0.72
        if item.bullets:
            bullet_box = slide.shapes.add_textbox(Inches(left), Inches(y), Inches(7.0 if item.code_blocks else 11.8), Inches(5.05))
            btf=bullet_box.text_frame; btf.clear(); btf.word_wrap=True
            btf.margin_left=Inches(0.12); btf.margin_right=Inches(0.08)
            for bi, bullet in enumerate(item.bullets):
                p=btf.paragraphs[0] if bi==0 else btf.add_paragraph()
                p.text=f"•  {bullet}"; p.font.name="Aptos"; p.font.size=Pt(22 if len(item.bullets)<=4 else 18); p.font.color.rgb=dark; p.space_after=Pt(15)
        if item.code_blocks:
            code = item.code_blocks[0]
            # Keep classroom readability; full source code is also included in separate files.
            lines = code.code.splitlines()
            fit_lines = lines[:18]
            display_code = "\n".join(fit_lines)
            if len(lines)>18: display_code += "\n… full source in 03_Extracted_Code"
            code_box=slide.shapes.add_textbox(Inches(7.95), Inches(1.55), Inches(4.75), Inches(4.8))
            ctf=code_box.text_frame; ctf.clear(); ctf.word_wrap=True; ctf.margin_left=Inches(0.18); ctf.margin_right=Inches(0.12); ctf.margin_top=Inches(0.15)
            cp=ctf.paragraphs[0]; cp.text=display_code; cp.font.name="Cascadia Code"; cp.font.size=Pt(14); cp.font.color.rgb=RGBColor(30,41,59)
            code_box.fill.solid(); code_box.fill.fore_color.rgb=RGBColor(243,246,250); code_box.line.color.rgb=RGBColor(210,218,230)
            if len(lines)>18 or max([len(l) for l in fit_lines] or [0])>52:
                overflow.append({"slide_id":f"SLIDE-{idx:02d}","issue":"Code may need inspection; slide shows a readable excerpt and full code is saved separately"})
        elif item.diagram_mermaid:
            # Create a native editable diagram instead of leaving a placeholder.
            labels = re.findall(r"\[([^\]]+)\]", item.diagram_mermaid)
            if not labels: labels=["Input","Process","Output"]
            labels=labels[:4]
            x_start=1.0; gap=2.5
            for di,label in enumerate(labels):
                x=x_start+di*gap
                shape=slide.shapes.add_shape(5, Inches(x), Inches(3.05), Inches(2.0), Inches(0.85))
                shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor(232,240,254); shape.line.color.rgb=accent
                shape.text_frame.text=label
                shape.text_frame.paragraphs[0].font.size=Pt(16); shape.text_frame.paragraphs[0].font.color.rgb=dark
                if di<len(labels)-1:
                    arrow=slide.shapes.add_shape(33, Inches(x+2.05), Inches(3.32), Inches(0.45), Inches(0.25))
                    arrow.fill.solid(); arrow.fill.fore_color.rgb=accent; arrow.line.fill.background()
        # Footer and number
        footer=slide.shapes.add_textbox(Inches(0.65), Inches(7.0), Inches(11.8), Inches(0.25))
        fp=footer.text_frame.paragraphs[0]; fp.text=f"{brand}  •  {course.title}"; fp.font.name="Aptos"; fp.font.size=Pt(9); fp.font.color.rgb=RGBColor(100,116,139)
        num=slide.shapes.add_textbox(Inches(12.35), Inches(7.0), Inches(0.5), Inches(0.25))
        np=num.text_frame.paragraphs[0]; np.text=str(idx); np.alignment=PP_ALIGN.RIGHT; np.font.size=Pt(9); np.font.color.rgb=accent
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return {"slide_count":len(course.slides),"render_engine":"python-pptx","overflow_warnings":overflow,"structural_overflow_count":len(overflow)}


def _make_pdf(course: CourseContent, path: Path, brand: str, primary: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name="SlideTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=24, leading=28, textColor=colors.HexColor(primary), spaceAfter=18))
    styles.add(ParagraphStyle(name="SlideBody", parent=styles["BodyText"], fontSize=16, leading=22, spaceAfter=10))
    styles.add(ParagraphStyle(name="SlideCode", parent=styles["Code"], fontName="Courier", fontSize=8.3, leading=10, backColor=colors.HexColor("#F3F6FA"), borderPadding=8))
    styles.add(ParagraphStyle(name="SlideCaption", parent=styles["BodyText"], fontName="Helvetica", fontSize=8, leading=10, textColor=colors.HexColor("#64748B")))
    doc=SimpleDocTemplate(str(path),pagesize=landscape(letter),rightMargin=0.6*inch,leftMargin=0.6*inch,topMargin=0.55*inch,bottomMargin=0.55*inch)
    story=[]
    for idx, slide in enumerate(course.slides):
        story.append(Paragraph(escape(slide.title),styles["SlideTitle"]))
        for bullet in slide.bullets:
            story.append(Paragraph("• &nbsp;"+escape(bullet),styles["SlideBody"]))
        if slide.code_blocks:
            code=slide.code_blocks[0].code[:2600]
            story.append(Spacer(1,8)); story.append(Preformatted(code,styles["SlideCode"]))
        elif slide.diagram_mermaid:
            nodes=re.findall(r"\[([^\]]+)\]",slide.diagram_mermaid)
            if nodes:
                t=Table([[escape(n) for n in nodes[:4]]],colWidths=[2.2*inch]*min(4,len(nodes)))
                t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#E8F0FE")),("BOX",(0,0),(-1,-1),1,colors.HexColor(primary)),("INNERGRID",(0,0),(-1,-1),1,colors.white),("ALIGN",(0,0),(-1,-1),"CENTER"),("TOPPADDING",(0,0),(-1,-1),18),("BOTTOMPADDING",(0,0),(-1,-1),18)])); story.append(t)
        story.append(Spacer(1,8)); story.append(Paragraph(f"{escape(brand)} · {idx+1}/{len(course.slides)}",styles["SlideCaption"]))
        if idx < len(course.slides)-1: story.append(PageBreak())
    doc.build(story)


def build_artifacts(root: Path, material_id: str, course: CourseContent, snippets: list[dict], report: dict, request_payload: dict, brand: str, primary: str) -> Path:
    out=root/material_id
    if out.exists(): shutil.rmtree(out)
    for folder in ["01_Presentation","02_Trainer_Notes","03_Extracted_Code","04_Verified_Code","05_Failed_Code","06_Exercises","07_Verification_Reports","08_Quiz_CheatSheet"]:
        (out/folder).mkdir(parents=True,exist_ok=True)
    _write(out/"course.json", course.model_dump_json(indent=2))
    _write(out/"request.json", json.dumps(request_payload,indent=2,ensure_ascii=False))
    _make_pptx(course,out/"01_Presentation/course_slides.pptx",brand,primary)
    _make_pdf(course,out/"01_Presentation/course_slides.pdf",brand,primary)
    # Notes files are always separate and have stable slide IDs.
    notes_parts=[]
    for i,slide in enumerate(course.slides,1):
        note=f"# SLIDE-{i:02d}: {slide.title}\n\n{slide.notes.strip()}\n"
        _write(out/f"02_Trainer_Notes/slide_{i:02d}_notes.md",note)
        notes_parts.append(note)
    _write(out/"02_Trainer_Notes/complete_trainer_notes.md","\n---\n\n".join(notes_parts))
    # Extracted code is saved independently, with provenance and status.
    _write(out/"03_Extracted_Code/snippet_manifest.json",json.dumps(snippets,indent=2,ensure_ascii=False))
    all_md=[]
    for snip in snippets:
        ext={"python":"py","java":"java"}.get(snip.get("language"),"txt")
        fname=f"{snip['snippet_id']}_{re.sub(r'[^A-Za-z0-9_-]+','_',snip.get('title','snippet'))[:40]}.{ext}"
        _write(out/f"03_Extracted_Code/{fname}",snip.get("code",""))
        all_md.append(f"## {snip['snippet_id']} — {snip.get('title','Snippet')}\n\nSource: `{json.dumps(snip.get('source_locations',[]),ensure_ascii=False)}`\n\n```{snip.get('language','text')}\n{snip.get('code','')}\n```\n")
        status=snip.get("status","PENDING")
        if status=="PASSED":
            safe_name=fname
            _write(out/f"04_Verified_Code/{safe_name}",snip.get("code",""))
        elif status in {"FAILED","BLOCKED","ERROR"}:
            _write(out/f"05_Failed_Code/{fname}",snip.get("code",""))
            _write(out/f"05_Failed_Code/{snip['snippet_id']}_error.md",f"# {snip['snippet_id']} — {status}\n\n{snip.get('error','No details returned')}\n")
        elif status=="ARTIFACT_ONLY":
            _write(out/f"06_Exercises/{snip['snippet_id']}_test_harness.{ext}",snip.get("code",""))
    _write(out/"03_Extracted_Code/all_snippets.md","\n---\n\n".join(all_md))
    # Exercise project bundles: reference / naive are stored separately, tests are preserved.
    for eidx, exercise in enumerate(course.exercises,1):
        base=out/f"06_Exercises/exercise_{eidx:02d}"
        ext="py" if exercise.language=="python" else "java"
        _write(base/"README.md",f"# {exercise.title}\n\n{exercise.prompt}\n\nStatus: see verification_report.json\n")
        _write(base/f"reference_solution/solution.{ext}",exercise.reference_solution)
        _write(base/f"naive_solution/solution.{ext}",exercise.naive_solution)
        test_filename = "test_solution.py" if exercise.language=="python" else "TestMain.java"
        _write(base/"tests"/test_filename,exercise.tests)
    _write(out/"08_Quiz_CheatSheet/quiz.md","# Quiz\n\n"+"\n\n".join(f"## Q{i}. {q.question}\n"+"\n".join(f"- {'✓ ' if oi==q.correct_index else ''}{opt}" for oi,opt in enumerate(q.options)) for i,q in enumerate(course.quiz,1)))
    _write(out/"08_Quiz_CheatSheet/quiz_answers_explanations.md","# Answers and explanations\n\n"+"\n\n".join(f"{i}. **{q.options[q.correct_index]}**\n\n{q.explanation}" for i,q in enumerate(course.quiz,1)))
    _write(out/"08_Quiz_CheatSheet/one_page_cheatsheet.md","# One-page cheat sheet\n\n"+"\n".join(f"- {x}" for x in course.cheat_sheet))
    # A compact PDF for the cheat sheet.
    cheat_pdf=out/"08_Quiz_CheatSheet/one_page_cheatsheet.pdf"
    cheat_styles=getSampleStyleSheet()
    cheat_doc=SimpleDocTemplate(str(cheat_pdf),pagesize=letter,rightMargin=0.7*inch,leftMargin=0.7*inch,topMargin=0.7*inch,bottomMargin=0.7*inch)
    cheat_story=[Paragraph(escape(course.title),cheat_styles["Title"]),Spacer(1,12)]
    for line in course.cheat_sheet: cheat_story.append(Paragraph("• &nbsp;"+escape(line),cheat_styles["BodyText"])); cheat_story.append(Spacer(1,8))
    cheat_doc.build(cheat_story)
    # Merge rendering statistics into report before writing.
    report["presentation"]={"pptx":"PASS","pdf":"PASS","slide_count":len(course.slides),"layout_checks":"structural approximation; browser/rendered-image inspection should be enabled for production"}
    _write(out/"07_Verification_Reports/verification_report.json",json.dumps(report,indent=2,ensure_ascii=False))
    _write(out/"07_Verification_Reports/verification_report.md",_report_markdown(report))
    _write(out/"07_Verification_Reports/confidence_score.json",json.dumps(report.get("confidence",{}),indent=2))
    logs=out/"07_Verification_Reports/execution_logs"; logs.mkdir(exist_ok=True)
    for s in snippets:
        if s.get("execution"):
            _write(logs/f"{s['snippet_id']}.log",f"status={s.get('status')}\nexit_code={s['execution'].get('exit_code')}\nstdout:\n{s['execution'].get('stdout','')}\nstderr:\n{s['execution'].get('stderr','')}\n")
    zip_path=root/f"{material_id}_course_pack.zip"
    if zip_path.exists(): zip_path.unlink()
    with zipfile.ZipFile(zip_path,"w",zipfile.ZIP_DEFLATED) as z:
        for file in out.rglob("*"):
            if file.is_file(): z.write(file,file.relative_to(out.parent))
    return zip_path


def _report_markdown(report: dict) -> str:
    lines=["# Verification Report", "", f"**Status:** {report.get('status','UNKNOWN')}", f"**Confidence:** {report.get('confidence',{}).get('score','N/A')}/100 ({report.get('confidence',{}).get('band','N/A')})", "", "## Summary", ""]
    for key,value in report.get("summary",{}).items(): lines.append(f"- **{key.replace('_',' ').title()}:** {value}")
    lines += ["", "## Snippet results", ""]
    for s in report.get("snippets",[]): lines.append(f"- `{s.get('snippet_id')}` — **{s.get('status')}** — {s.get('title')} — {s.get('error','')}")
    return "\n".join(lines)+"\n"


def refresh_report_and_zip(root: Path, material_id: str, report: dict) -> Path:
    """Rewrite the evidence report after rendered-PDF inspection, then rebuild the ZIP."""
    out=root/material_id
    report_dir=out/"07_Verification_Reports"
    report_dir.mkdir(parents=True,exist_ok=True)
    _write(report_dir/"verification_report.json",json.dumps(report,indent=2,ensure_ascii=False))
    _write(report_dir/"verification_report.md",_report_markdown(report))
    _write(report_dir/"confidence_score.json",json.dumps(report.get("confidence",{}),indent=2,ensure_ascii=False))
    zip_path=root/f"{material_id}_course_pack.zip"
    if zip_path.exists(): zip_path.unlink()
    with zipfile.ZipFile(zip_path,"w",zipfile.ZIP_DEFLATED) as z:
        for file in out.rglob("*"):
            if file.is_file(): z.write(file,file.relative_to(out.parent))
    return zip_path
