import json
import zipfile
from pathlib import Path

from pptx import Presentation

from app.extraction import extract_snippets
from app.generator import _demo_content
from app.models import GenerateRequest
from app.packager import build_artifacts


def test_course_pack_includes_pptx_speaker_notes_and_separate_files(tmp_path: Path):
    request = GenerateRequest(topic="Python functions", programming_language="python")
    course = _demo_content(request)
    snippets = extract_snippets(course)
    report = {"status": "BLOCKED", "confidence": {"score": 34, "band": "LOW"}, "summary": {}}

    archive = build_artifacts(
        tmp_path,
        "mat_pack_test",
        course,
        snippets,
        report,
        request.model_dump(),
        request.brand_name,
        request.brand_primary_color,
    )

    presentation = Presentation(str(tmp_path / "mat_pack_test/01_Presentation/course_slides.pptx"))
    assert len(presentation.slides) == len(course.slides)
    assert "Explain why this topic matters" in presentation.slides[0].notes_slide.notes_text_frame.text
    assert (tmp_path / "mat_pack_test/02_Trainer_Notes/slide_01_notes.md").exists()
    manifest = json.loads((tmp_path / "mat_pack_test/03_Extracted_Code/snippet_manifest.json").read_text())
    assert len(manifest) == len(snippets)
    with zipfile.ZipFile(archive) as zf:
        names = set(zf.namelist())
    assert "mat_pack_test/01_Presentation/course_slides.pptx" in names
    assert "mat_pack_test/02_Trainer_Notes/slide_01_notes.md" in names
    assert "mat_pack_test/03_Extracted_Code/snippet_manifest.json" in names
